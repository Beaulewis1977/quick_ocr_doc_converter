#!/usr/bin/env python3
"""
Document to Markdown Converter - Desktop GUI
A simple desktop application for converting documents to Markdown format
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import threading
import os
from pathlib import Path
import sys

class DocumentConverterGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Document to Markdown Converter")
        self.root.geometry("600x500")
        self.root.minsize(500, 400)
        
        # Variables
        self.input_path = tk.StringVar()
        self.output_path = tk.StringVar()
        self.progress_var = tk.DoubleVar()
        self.status_var = tk.StringVar(value="Ready to convert documents")
        
        # Set default output to Desktop/markdown_output
        desktop = Path.home() / "Desktop"
        default_output = desktop / "markdown_output"
        self.output_path.set(str(default_output))
        
        self.setup_ui()
        
    def setup_ui(self):
        """Set up the user interface"""
        # Main frame
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        
        # Configure grid weights
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(0, weight=1)
        main_frame.columnconfigure(1, weight=1)
        
        # Title
        title_label = ttk.Label(main_frame, text="Document to Markdown Converter", 
                               font=('Arial', 16, 'bold'))
        title_label.grid(row=0, column=0, columnspan=3, pady=(0, 20))
        
        # Input folder selection
        ttk.Label(main_frame, text="Source Folder:", font=('Arial', 10, 'bold')).grid(
            row=1, column=0, sticky=tk.W, pady=(0, 5))
        
        input_frame = ttk.Frame(main_frame)
        input_frame.grid(row=2, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(0, 10))
        input_frame.columnconfigure(0, weight=1)
        
        self.input_entry = ttk.Entry(input_frame, textvariable=self.input_path, 
                                    font=('Arial', 9))
        self.input_entry.grid(row=0, column=0, sticky=(tk.W, tk.E), padx=(0, 5))
        
        ttk.Button(input_frame, text="Browse", 
                  command=self.browse_input_folder).grid(row=0, column=1)
        
        # Output folder selection
        ttk.Label(main_frame, text="Output Folder:", font=('Arial', 10, 'bold')).grid(
            row=3, column=0, sticky=tk.W, pady=(0, 5))
        
        output_frame = ttk.Frame(main_frame)
        output_frame.grid(row=4, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(0, 10))
        output_frame.columnconfigure(0, weight=1)
        
        self.output_entry = ttk.Entry(output_frame, textvariable=self.output_path, 
                                     font=('Arial', 9))
        self.output_entry.grid(row=0, column=0, sticky=(tk.W, tk.E), padx=(0, 5))
        
        ttk.Button(output_frame, text="Browse", 
                  command=self.browse_output_folder).grid(row=0, column=1)
        
        # Options frame
        options_frame = ttk.LabelFrame(main_frame, text="Options", padding="10")
        options_frame.grid(row=5, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(0, 15))
        
        self.preserve_structure = tk.BooleanVar(value=True)
        ttk.Checkbutton(options_frame, text="Preserve folder structure", 
                       variable=self.preserve_structure).grid(row=0, column=0, sticky=tk.W)
        
        self.overwrite_existing = tk.BooleanVar(value=False)
        ttk.Checkbutton(options_frame, text="Overwrite existing files", 
                       variable=self.overwrite_existing).grid(row=1, column=0, sticky=tk.W)
        
        # Convert button
        self.convert_button = ttk.Button(main_frame, text="Convert Documents", 
                                        command=self.start_conversion,
                                        style='Accent.TButton')
        self.convert_button.grid(row=6, column=0, columnspan=3, pady=(0, 15))
        
        # Progress bar
        ttk.Label(main_frame, text="Progress:", font=('Arial', 10, 'bold')).grid(
            row=7, column=0, sticky=tk.W)
        
        self.progress_bar = ttk.Progressbar(main_frame, variable=self.progress_var, 
                                           maximum=100)
        self.progress_bar.grid(row=8, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(0, 10))
        
        # Status label
        self.status_label = ttk.Label(main_frame, textvariable=self.status_var, 
                                     font=('Arial', 9), foreground='blue')
        self.status_label.grid(row=9, column=0, columnspan=3, sticky=tk.W)
        
        # Results text area
        ttk.Label(main_frame, text="Results:", font=('Arial', 10, 'bold')).grid(
            row=10, column=0, sticky=tk.W, pady=(15, 5))
        
        # Create text widget with scrollbar
        text_frame = ttk.Frame(main_frame)
        text_frame.grid(row=11, column=0, columnspan=3, sticky=(tk.W, tk.E, tk.N, tk.S), pady=(0, 10))
        text_frame.columnconfigure(0, weight=1)
        text_frame.rowconfigure(0, weight=1)
        main_frame.rowconfigure(11, weight=1)
        
        self.results_text = tk.Text(text_frame, height=8, wrap=tk.WORD, font=('Consolas', 9))
        scrollbar = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self.results_text.yview)
        self.results_text.configure(yscrollcommand=scrollbar.set)
        
        self.results_text.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        scrollbar.grid(row=0, column=1, sticky=(tk.N, tk.S))
        
        # Drag and drop setup
        self.setup_drag_drop()
        
    def setup_drag_drop(self):
        """Set up drag and drop functionality"""
        try:
            from tkinterdnd2 import TkinterDnD, DND_FILES
            # If tkinterdnd2 is available, enable drag and drop
            self.root.drop_target_register(DND_FILES)
            self.root.dnd_bind('<<Drop>>', self.on_drop)
            
            # Add instruction label
            instruction_text = "💡 Tip: You can drag and drop folders directly onto this window!"
            ttk.Label(self.root, text=instruction_text, font=('Arial', 8), 
                     foreground='gray').grid(row=1, column=0, pady=(0, 5))
        except ImportError:
            # tkinterdnd2 not available, show regular instructions
            instruction_text = "💡 Tip: Use the Browse buttons to select folders"
            ttk.Label(self.root, text=instruction_text, font=('Arial', 8), 
                     foreground='gray').grid(row=1, column=0, pady=(0, 5))
    
    def on_drop(self, event):
        """Handle drag and drop events"""
        files = self.root.tk.splitlist(event.data)
        if files:
            # Use the first dropped item as input folder
            dropped_path = files[0]
            if os.path.isdir(dropped_path):
                self.input_path.set(dropped_path)
                self.log_message(f"📁 Folder dropped: {dropped_path}")
            else:
                # If it's a file, use its parent directory
                parent_dir = os.path.dirname(dropped_path)
                self.input_path.set(parent_dir)
                self.log_message(f"📁 File dropped, using parent folder: {parent_dir}")
    
    def browse_input_folder(self):
        """Browse for input folder"""
        folder = filedialog.askdirectory(title="Select folder containing documents to convert")
        if folder:
            self.input_path.set(folder)
    
    def browse_output_folder(self):
        """Browse for output folder"""
        folder = filedialog.askdirectory(title="Select output folder for markdown files")
        if folder:
            self.output_path.set(folder)
    
    def validate_inputs(self):
        """Validate user inputs before conversion"""
        if not self.input_path.get():
            messagebox.showerror("Error", "Please select an input folder")
            return False
        
        if not os.path.exists(self.input_path.get()):
            messagebox.showerror("Error", "Input folder does not exist")
            return False
        
        if not self.output_path.get():
            messagebox.showerror("Error", "Please select an output folder")
            return False
        
        return True
    
    def start_conversion(self):
        """Start the conversion process in a separate thread"""
        if not self.validate_inputs():
            return
        
        # Disable the convert button
        self.convert_button.config(state='disabled')
        self.progress_var.set(0)
        self.results_text.delete(1.0, tk.END)
        
        # Start conversion in a separate thread
        thread = threading.Thread(target=self.convert_documents, daemon=True)
        thread.start()
    
    def convert_documents(self):
        """Convert documents (runs in separate thread)"""
        try:
            # Import conversion modules
            self.update_status("Installing required packages...")
            self.install_requirements()
            
            self.update_status("Scanning for documents...")
            
            # Import conversion functions
            from docx import Document
            import PyPDF2
            
            input_dir = Path(self.input_path.get())
            output_dir = Path(self.output_path.get())
            
            # Find all files to convert
            file_patterns = ['*.docx', '*.pdf', '*.txt']
            files_to_convert = []
            
            for pattern in file_patterns:
                files_to_convert.extend(input_dir.rglob(pattern))
            
            # Filter out temporary files
            files_to_convert = [f for f in files_to_convert if not f.name.startswith('~$')]
            
            total_files = len(files_to_convert)
            
            if total_files == 0:
                self.log_message("❌ No convertible files found (looking for .docx, .pdf, .txt)")
                self.update_status("No files to convert")
                self.convert_button.config(state='normal')
                return
            
            self.log_message(f"📊 Found {total_files} files to convert")
            self.update_status(f"Converting {total_files} files...")
            
            # Create output directory
            output_dir.mkdir(parents=True, exist_ok=True)
            
            successful = 0
            failed = 0
            
            for i, file_path in enumerate(files_to_convert):
                try:
                    # Calculate relative path for preserving structure
                    if self.preserve_structure.get():
                        rel_path = file_path.relative_to(input_dir)
                        output_file_path = output_dir / rel_path.with_suffix('.md')
                    else:
                        output_file_path = output_dir / f"{file_path.stem}.md"
                    
                    # Create output directory if needed
                    output_file_path.parent.mkdir(parents=True, exist_ok=True)
                    
                    # Skip if file exists and not overwriting
                    if output_file_path.exists() and not self.overwrite_existing.get():
                        self.log_message(f"⏭️  Skipped (exists): {file_path.name}")
                        continue
                    
                    # Convert based on file type
                    if file_path.suffix.lower() == '.docx':
                        content = self.convert_docx_to_markdown(file_path)
                    elif file_path.suffix.lower() == '.pdf':
                        content = self.convert_pdf_to_markdown(file_path)
                    elif file_path.suffix.lower() == '.txt':
                        content = self.convert_txt_to_markdown(file_path)
                    else:
                        continue
                    
                    # Save markdown file
                    with open(output_file_path, 'w', encoding='utf-8') as f:
                        f.write(content)
                    
                    successful += 1
                    self.log_message(f"✅ {file_path.name}")
                    
                except Exception as e:
                    failed += 1
                    self.log_message(f"❌ {file_path.name}: {str(e)}")
                
                # Update progress
                progress = ((i + 1) / total_files) * 100
                self.progress_var.set(progress)
            
            # Final results
            self.log_message(f"\n🎉 Conversion complete!")
            self.log_message(f"✅ Successful: {successful}")
            self.log_message(f"❌ Failed: {failed}")
            self.log_message(f"📁 Output saved to: {output_dir}")
            
            self.update_status(f"Complete! {successful} files converted")
            
            # Show completion message
            messagebox.showinfo("Conversion Complete", 
                              f"Successfully converted {successful} files!\n"
                              f"Failed: {failed}\n"
                              f"Output saved to: {output_dir}")
            
        except Exception as e:
            self.log_message(f"❌ Error during conversion: {str(e)}")
            self.update_status("Conversion failed")
            messagebox.showerror("Error", f"Conversion failed: {str(e)}")
        
        finally:
            # Re-enable the convert button
            self.convert_button.config(state='normal')
    
    def install_requirements(self):
        """Install required packages"""
        required_packages = ['python-docx', 'PyPDF2']
        
        for package in required_packages:
            try:
                __import__(package.replace('-', '_'))
            except ImportError:
                self.log_message(f"📦 Installing {package}...")
                os.system(f'pip install {package}')
    
    def convert_docx_to_markdown(self, file_path):
        """Convert DOCX file to Markdown"""
        from docx import Document
        
        doc = Document(file_path)
        markdown_content = []
        
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                # Simple markdown conversion
                if paragraph.style.name.startswith('Heading'):
                    level = int(paragraph.style.name.split()[-1]) if paragraph.style.name.split()[-1].isdigit() else 1
                    markdown_content.append(f"{'#' * level} {text}")
                else:
                    markdown_content.append(text)
                markdown_content.append("")  # Add blank line
        
        return "\n".join(markdown_content)
    
    def convert_pdf_to_markdown(self, file_path):
        """Convert PDF file to Markdown"""
        import PyPDF2
        
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            text_content = []
            
            for page in pdf_reader.pages:
                text_content.append(page.extract_text())
        
        # Join all pages and clean up
        full_text = "\n\n".join(text_content)
        # Basic cleanup
        lines = [line.strip() for line in full_text.split('\n') if line.strip()]
        return "\n\n".join(lines)
    
    def convert_txt_to_markdown(self, file_path):
        """Convert TXT file to Markdown"""
        encodings = ['utf-8', 'latin-1', 'cp1252', 'iso-8859-1']
        
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as file:
                    content = file.read()
                return content
            except UnicodeDecodeError:
                continue
        
        raise Exception(f"Could not decode file with any of the attempted encodings: {encodings}")
    
    def update_status(self, message):
        """Update status label thread-safely"""
        self.root.after(0, lambda: self.status_var.set(message))
    
    def log_message(self, message):
        """Add message to results text area thread-safely"""
        def _add_message():
            self.results_text.insert(tk.END, message + "\n")
            self.results_text.see(tk.END)
        
        self.root.after(0, _add_message)

def main():
    """Main application entry point"""
    try:
        # Try to use tkinterdnd2 for drag and drop
        from tkinterdnd2 import TkinterDnD
        root = TkinterDnD.Tk()
    except ImportError:
        # Fall back to regular tkinter
        root = tk.Tk()
    
    app = DocumentConverterGUI(root)
    
    # Center the window
    root.update_idletasks()
    x = (root.winfo_screenwidth() // 2) - (root.winfo_width() // 2)
    y = (root.winfo_screenheight() // 2) - (root.winfo_height() // 2)
    root.geometry(f"+{x}+{y}")
    
    root.mainloop()

if __name__ == "__main__":
    main() 